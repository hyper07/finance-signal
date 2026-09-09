"""Build and validate the double-blind editable FRL manuscript."""
from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path

import pypandoc
from docx import Document
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parent.parent
PAPER = REPO / "research" / "papers" / "frl_submission"
SOURCE = PAPER / "paper.md"
ANONYMOUS_MD = PAPER / "Manuscript_Anonymous.md"
ANONYMOUS_DOCX = PAPER / "Manuscript_Anonymous.docx"
PRIVATE = PAPER / "private"

IDENTIFIERS = (
    "Kibaek",
    "Kiok",
    "Danielle",
    "dotori.ai",
    "hyper07",
)


def anonymize(text: str) -> str:
    text = re.sub(
        r"^\*Kibaek Kim, Kiok Kim, and Danielle Ahn — dotori\.ai\*\n\n",
        "",
        text,
        flags=re.MULTILINE,
    )
    text = text.replace(
        "The authors operate the audited service.",
        "Members of the research team operate the audited service.",
    )
    text = text.replace(
        "Rights-safe aggregate results and analysis code are available at "
        "<https://github.com/hyper07/finance-signal>.",
        "Rights-safe aggregate results and analysis code will be released in a "
        "public repository after review; anonymized materials are available to "
        "the editor on request.",
    )
    text = text.replace(
        "The authors operate the forecasting service evaluated in this study.",
        "Members of the research team operate the forecasting service evaluated "
        "in this study.",
    )
    return text


def apply_word_formatting(path: Path) -> None:
    document = Document(path)
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 2

    for paragraph in document.paragraphs:
        paragraph.paragraph_format.line_spacing = 2
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 2

    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = (
        "When information breaks the historical pattern: preliminary evidence "
        "from a deployed Bitcoin forecaster"
    )
    properties.subject = "Anonymous manuscript"
    properties.comments = ""
    document.save(path)


def validate_anonymity(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        searchable = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )
    found = [token for token in IDENTIFIERS if token.lower() in searchable.lower()]
    if found:
        raise ValueError(f"Author identifiers remain in {path.name}: {found}")


def build_title_page(email: str, address: str, affiliation: str) -> Path:
    PRIVATE.mkdir(exist_ok=True)
    path = PRIVATE / "Title_Page.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = document.add_heading(
        "When information breaks the historical pattern: preliminary evidence "
        "from a deployed Bitcoin forecaster",
        level=0,
    )
    title.alignment = 1

    authors = document.add_paragraph()
    authors.alignment = 1
    authors.add_run("Kibaek Kim")
    marker = authors.add_run("a,*")
    marker.font.superscript = True
    authors.add_run("; Kiok Kim")
    marker = authors.add_run("a")
    marker.font.superscript = True
    authors.add_run("; Danielle Ahn")
    marker = authors.add_run("a")
    marker.font.superscript = True

    affiliation_line = document.add_paragraph()
    affiliation_line.alignment = 1
    marker = affiliation_line.add_run("a")
    marker.font.superscript = True
    affiliation_line.add_run(f" {affiliation}, {address}")

    corresponding = document.add_paragraph()
    corresponding.alignment = 1
    marker = corresponding.add_run("*")
    marker.font.superscript = True
    corresponding.add_run(
        f" Corresponding author: Kibaek Kim, {email}; {address}"
    )

    document.add_heading("Declaration of competing interest", level=1)
    document.add_paragraph(
        "The authors operate the forecasting service evaluated in this study. "
        "No investment recommendation is made."
    )

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    properties = document.core_properties
    properties.author = "Kibaek Kim, Kiok Kim, Danielle Ahn"
    properties.last_modified_by = ""
    properties.title = title.text
    properties.subject = "Title page"
    document.save(path)
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--corresponding-email")
    parser.add_argument("--postal-address")
    parser.add_argument("--affiliation", default="dotori.ai")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    anonymous = anonymize(SOURCE.read_text())
    for token in IDENTIFIERS:
        if token.lower() in anonymous.lower():
            raise ValueError(f"Author identifier remains in Markdown: {token}")
    ANONYMOUS_MD.write_text(anonymous)
    pypandoc.convert_file(
        str(ANONYMOUS_MD),
        "docx",
        outputfile=str(ANONYMOUS_DOCX),
        format="markdown+tex_math_dollars+pipe_tables",
        extra_args=[f"--resource-path={PAPER}"],
    )
    apply_word_formatting(ANONYMOUS_DOCX)
    validate_anonymity(ANONYMOUS_DOCX)
    print("written:", ANONYMOUS_DOCX)
    provided = (args.corresponding_email, args.postal_address)
    if any(provided) and not all(provided):
        raise ValueError("Provide both --corresponding-email and --postal-address")
    if all(provided):
        title_page = build_title_page(
            args.corresponding_email,
            args.postal_address,
            args.affiliation,
        )
        print("written:", title_page)


if __name__ == "__main__":
    main()
